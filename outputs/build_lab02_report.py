from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path.cwd()
TEMPLATE = next(Path('..').glob('*.docx'))
OUTPUT = ROOT / 'submissions/lab02/01_实验报告/实验报告_何嘉欣_2410311411.docx'
DOCX_OUTPUTS = ROOT / 'outputs/lab02/docx'

doc = Document(str(TEMPLATE))
doc.core_properties.title = '云计算与大数据处理实验报告：实验二 数据湖与质量契约'
doc.core_properties.author = '何嘉欣'
doc.core_properties.subject = '实验二 数据湖与质量契约'

cover_values = {
    '实验名称': '实验二　数据湖与质量契约',
    '实验日期': '2026年10月26日',
    '专业': '计算机科学与技术',
    '班级': '24计科4',
    '学号': '2410311411',
    '姓名': '何嘉欣',
    '指导教师': '刘锦行',
}
for row in doc.tables[0].rows:
    key = row.cells[0].text.strip()
    if key in cover_values:
        cell = row.cells[1]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(cover_values[key])
        run.font.name = '宋体'
        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
        run.font.size = Pt(12)
        run.font.bold = False
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

normal = doc.styles['Normal']
normal.font.name = '宋体'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(4)


def set_run_font(run, name='宋体', size=10.5, bold=False, color='000000'):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def para(text='', bold=False, size=10.5, align=None, name='宋体', space_after=4, first_line=True):
    p = doc.add_paragraph()
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(space_after)
    if first_line:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    if text:
        run = p.add_run(text)
        set_run_font(run, name, size, bold)
    return p


def heading(text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, '黑体', {1: 16, 2: 13, 3: 11.5}[level], True)
    return p


def bullet(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.42)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run('• ' + text)
    set_run_font(run, '宋体', 10.5)
    return p


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.insert_element_before(shd, 'w:noWrap', 'w:tcMar', 'w:textDirection',
                                'w:tcFitText', 'w:vAlign', 'w:hideMark')


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '808080')
        borders.append(el)
    tbl_pr.insert_element_before(borders, 'w:shd', 'w:tblLayout', 'w:tblCellMar',
                                 'w:tblLook', 'w:tblCaption', 'w:tblDescription')


def add_table(headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Normal Table'
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(text))
        set_run_font(run, '黑体', font_size, True)
        shade(cell, 'D9E2F3')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            set_run_font(run, '宋体', font_size)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Inches(width)
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def code_block(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.0
    lines = text.splitlines()
    for i, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, 'Consolas', 8.5)
        if i < len(lines) - 1:
            run.add_break()


def image_with_caption(filename, caption, width=5.45):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(ROOT / 'submissions/lab02/evidence/截图' / filename), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    run = cap.add_run(caption)
    set_run_font(run, '宋体', 9)


while doc.paragraphs and not doc.paragraphs[-1].text.strip():
    element = doc.paragraphs[-1]._element
    element.getparent().remove(element)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.page_break_before = True
run = p.add_run('实验二　数据湖与质量契约')
set_run_font(run, '黑体', 18, True)
p.paragraph_format.space_after = Pt(12)

heading('一、实验目的', 1)
para('本实验把实验一“字节能够保存并往返”的结论推进到“数据可以被可靠分析”。以 20,102 行合成服务事件为输入，建立 Bronze、Silver、Gold 三层：原始层保留全部记录，Silver 按明确契约校验、隔离和去重后写成 Parquet，Gold 用固定 SQL 口径按日期与服务聚合。')
para('实验重点是证明清洗过程的守恒关系，区分非法、重复和同 ID 冲突，理解严格类型校验、UTC 时间、Parquet 行组、列裁剪和谓词下推，并通过个人阈值 2000→1200 的改动说明技术校验与业务政策的边界。')

heading('二、实验环境与前置数据', 1)
add_table(['项目', '实测值'], [
    ['操作系统', 'Windows 11'],
    ['Python', '3.12.13'],
    ['项目解释器', r'D:\云计算与大数据处理\reference-BDP\.venv\Scripts\python.exe'],
    ['工作目录', r'D:\云计算与大数据处理\reference-BDP'],
    ['DuckDB / PyArrow', '1.5.5 / 25.0.1'],
    ['原始文件', 'data/raw/events.jsonl'],
    ['原始规模', '20,102 行，3,892,184 字节'],
    ['原始 SHA-256', '0d93bcd25e28047f804238c220bdb5b1d2dadf328647c07de85ba0104e8e2854'],
], widths=[1.15, 5.15], font_size=8.5)
para('原始数据由实验一继承，随机种子为 42；raw.json 记录合法独特 20,000、重复 100、非法 2。实验二的基线运行 run_id 为 e6a3a1ca90d1478688541f770a936ddf，stage 为 verify。清洗前后没有重新生成原始文件。')
image_with_caption('01-baseline-quality.png', '图 1　实验二基线质量、守恒关系与服务汇总（学号、实验号、run_id 与结果同屏）', 4.25)

heading('三、数据契约与处理规则', 1)
add_table(['字段', '契约', '拒绝原因示例'], [
    ['event_id', '非空字符串，事件主键', 'invalid_event_id'],
    ['event_time', '可解析的 ISO 时间，且 UTC 偏移必须为 0', 'invalid_event_time / non_utc_time'],
    ['service', 'assistant、search、batch', 'invalid_category'],
    ['model', 'small、medium', 'invalid_category'],
    ['input_tokens', '严格 int，且大于等于 0', 'invalid_input_tokens'],
    ['output_tokens', '严格 int，且大于等于 0', 'invalid_output_tokens'],
    ['latency_ms', '严格 int，且大于等于 0', 'invalid_latency_ms'],
    ['status', 'ok、error', 'invalid_status'],
], widths=[1.2, 3.3, 1.8], font_size=8.5)
para('字段先检查是否存在，再检查类型、范围和枚举，最后解析时间。type(x) is int 是严格判断：布尔值虽然继承 int，但不作为 1 Token 接收；数字字符串也不隐式转换，以免掩盖上游接口错误。UTC 时间先解析，再统一规范化为 datetime.isoformat() 后参与去重。')
para('event_id 相同且其余内容完全一致时计为重复；event_id 相同但内容不同时视为冲突并停止发布。不能默认“最后一条正确”，必须先定义更正、版本或纠错事件政策。拒绝记录写入 quarantine，保留原始行号、原因和原值，不静默删除。')

heading('四、独立手算：六行小表', 1)
add_table(['到达顺序', 'event_id', 'input_tokens', '特别说明', '处理结果'], [
    ['1', 'a', '10', '合法', '保留'],
    ['2', 'b', '20', '合法', '保留'],
    ['3', 'a', '10', '与第一行完全相同', '重复'],
    ['4', 'c', '-1', '非法负数', '隔离：invalid_input_tokens'],
    ['5', 'd', '40', '时间为 not-a-time', '隔离：invalid_event_time'],
    ['6', 'e', '30', '合法', '保留'],
], widths=[0.8, 0.75, 1.0, 1.75, 2.1], font_size=8.2)
para('手算结果：合法独特 3 条（a、b、e），重复 1 条，隔离 2 条，保留事件的 input_tokens 总量为 10+20+30=60。数量守恒为 3+1+2=6。')
para('若第 3 行改为 event_id=a、input_tokens=99，则它与第 1 行 ID 相同而内容不同，不能再计为普通重复，也不能简单保留最后一条。正确做法是停止发布并依据版本号、更新时间或显式纠错事件确定哪条代表业务事实；本实验的共同数据选择报冲突，避免在规则未定义时静默覆盖。')

heading('五、基线质量与隔离结果', 1)
add_table(['指标', '实测值', '说明'], [
    ['原始条数', '20,102', 'JSONL 总行数'],
    ['合法独特', '20,000', '写入 Silver 的主键唯一记录'],
    ['重复', '100', 'ID 与全部字段完全相同'],
    ['拒绝', '2', '进入隔离区并记录原因'],
    ['守恒', '20,102', '20000+100+2'],
    ['输入 Token 总量', '18,015,640', '仅统计合法独特记录'],
], widths=[1.55, 1.65, 3.1], font_size=8.8)
add_table(['原始行号', 'event_id', '拒绝原因', '原值'], [
    ['20,101', 'bad-negative', 'invalid_input_tokens', 'input_tokens=-1'],
    ['20,102', 'bad-time', 'invalid_event_time', 'event_time=not-a-time'],
], widths=[1.0, 1.35, 1.95, 2.0], font_size=8.5)
para('两条拒绝都不是流水线失败，而是契约按预期隔离坏记录。第一条类型为整数但范围非法；第二条字段存在但无法解析为 ISO 时间。保留原值使“哪一行、为什么拒绝、原值是什么”能够追溯。')
add_table(['service', '请求数', '输入 Token', '输出 Token', '延迟总和(ms)', '错误请求'], [
    ['assistant', '14,008', '12,631,850', '4,310,029', '28,448,309', '543'],
    ['batch', '2,023', '1,816,567', '616,515', '4,079,006', '78'],
    ['search', '3,969', '3,567,223', '1,219,318', '8,137,035', '179'],
], widths=[1.0, 0.9, 1.2, 1.1, 1.35, 0.75], font_size=8.2)
para('服务汇总的请求数合计为 20,000，输入 Token 合计为 18,015,640，输出 Token 合计为 6,145,862。Gold 统计合法日志中的全部请求，status=error 的 800 条请求仍计入请求数和资源口径；若只对成功请求计费，必须另行定义并解释口径。')

heading('六、Parquet 存储与 SQL 聚合', 1)
para('Silver 文件使用 Zstd 压缩，row_group_size=4096。实际写入 20,000 行、5 个行组，schema 含 string 类型的事件与类别字段、int64 类型的三个计数字段。原始 JSONL 为 3,892,184 字节，Silver Parquet 为 224,693 字节，约缩小到原来的 5.8%；但这不是纯格式性能对照，因为 Silver 已去掉坏行和重复，JSONL 是行式文本，Parquet 是列式压缩文件。')
image_with_caption('03-parquet-rowgroups.png', '图 2　Parquet schema、20,000 行与 5 个行组实测（run_id f6457389…814f53）', 5.1)
para('查询计划显示 READ_PARQUET 的投影只有 input_tokens，过滤条件为 service=\'search\'，读取 1 个文件和 3,969 行。这说明查询只投影需要的列，并把过滤条件下推到扫描阶段；不能据此宣称已经发生目录分区裁剪。本章只有一个 Parquet 文件，没有按日期或服务建立目录分区；行组是文件内部的水平批次，目录分区是按路径键值拆分数据，两者不是同一概念。')

code_block('''SELECT CAST(CAST(event_time AS TIMESTAMP) AS DATE)::VARCHAR AS day, service,
       COUNT(*) AS requests, SUM(input_tokens)::BIGINT, SUM(output_tokens)::BIGINT
FROM events
GROUP BY day, service
ORDER BY day, service;''')

add_table(['day', 'service', 'requests', 'input_tokens', 'output_tokens'], [
    ['2026-03-01', 'assistant', '12,142', '10,943,399', '3,743,595'],
    ['2026-03-01', 'batch', '1,722', '1,542,888', '526,892'],
    ['2026-03-01', 'search', '3,416', '3,066,824', '1,043,954'],
    ['2026-03-02', 'assistant', '1,866', '1,688,451', '566,434'],
    ['2026-03-02', 'batch', '301', '273,679', '89,623'],
    ['2026-03-02', 'search', '553', '500,399', '175,364'],
], widths=[1.25, 1.0, 0.9, 1.3, 1.3], font_size=8.3)

image_with_caption('02-gold-totals.png', '图 3　Silver 与 Gold 请求数和 Token 总量守恒（run_id 032c60ce…5b2af）', 5.35)
para('核对结果为 silver_totals=(20000, 18015640, 6145862)，gold_totals=(20000, 18015640, 6145862)，totals_match=true。分组粒度不同但总量一致，说明聚合没有无故丢账。')

heading('七、个人改动：输入 Token 上限 2000→1200', 1)
para('修改前先预测：课程合法数据的 input_tokens 最大不超过 1800，因此 2000 上限不会拒绝任何已合法事件；把上限改为 1200 后，超过 1200 的合法请求会被附加业务政策拒绝，合法独特数和输入 Token 总量都应下降。由于脚本先校验再去重，原本可计入重复的高 Token 行也会先进入拒绝计数，因此重复数可能从 100 降到更低。')
code_block('''- MAX_INPUT_TOKENS = 2000
 MAX_INPUT_TOKENS = 1200''')
add_table(['规则', '原始条数', '合法独特', '重复', '拒绝', '输入 Token 总量'], [
    ['基础课程契约', '20,102', '20,000', '100', '2', '18,015,640'],
    ['附加上限 2000', '20,102', '20,000', '100', '2', '18,015,640'],
    ['附加上限 1200', '20,102', '13,405', '67', '6,630', '8,110,488'],
], widths=[1.45, 0.9, 0.95, 0.7, 0.75, 1.4], font_size=8.2)
para('1200 规则的拒绝构成为 policy_input_limit=6,628、invalid_input_tokens=1、invalid_event_time=1；守恒为 13,405+67+6,630=20,102。与 2000 相比，合法独特减少 6,595 条，重复减少 33 条，说明有 33 条重复记录也因超过阈值而先被拒绝。输入 Token 总量减少 9,905,152，约占原来的 55.0%。')

image_with_caption('05-policy-2000.png', '图 4　个人政策 2000 上限：20,000 条合法独特，拒绝仅来自两类坏记录', 5.15)
image_with_caption('06-policy-1200.png', '图 5　个人政策 1200 上限：13,405 条合法独特，6,628 条因业务上限被拒绝', 5.15)
para('技术校验回答“记录是否合法”，业务政策回答“某类合法请求是否属于本次分析范围”。1200 上限不会把负数或坏时间变合法，也不会修复重复，它只是把“技术上合法但输入量高于业务阈值”的请求排除。该阈值适合成本控制或小请求分析，但会丢失大请求，不能直接替代课程共同契约，也不能覆盖实验三、四使用的 Silver。')

heading('八、契约探针、预期失败与恢复', 1)
add_table(['探针', '输入', '实测结果', '解释'], [
    ['布尔值', 'True', 'invalid_input_tokens', 'bool 是 int 子类，但契约要求精确 int'],
    ['数字字符串', '100', 'invalid_input_tokens', '不做隐式字符串转整数'],
    ['整数', '100', 'accepted', '类型和范围均符合契约'],
    ['负数', '-1', 'invalid_input_tokens', '范围小于 0'],
    ['无时区时间', '2026-03-01T12:00:00', 'non_utc_time', '缺少 UTC 偏移'],
], widths=[1.0, 1.3, 1.45, 2.55], font_size=8.3)
para('布尔值和字符串都得到 invalid_input_tokens，但原因不同：True 的问题在于类型不是精确 int，即使它可被 Python 当作整数参与运算；100 的问题在于值是字符串，必须由上游明确修正类型或定义转换规则。二者若被静默接受，都会掩盖接口契约错误。')
image_with_caption('07-contract-probes.png', '图 6　布尔值、数字字符串、整数、负数与无时区时间的契约探针（run_id 5be251e8…b99ea7）', 5.25)
para('预期失败运行时，把 100 假定为可接受输入，实际返回 exit_code=1、reason=invalid_input_tokens，stderr 为 EXPECTED_ACCEPT_BUT_GOT_invalid_input_tokens。恢复时把样本改成整数 100，运行返回 exit_code=0、reason=accepted。修复的是测试样本的类型，不是放宽全局校验。')
image_with_caption('08-expected-failure.png', '图 7　数字字符串被拒绝的预期失败（run_id b54b88af…04151b）', 4.35)
image_with_caption('09-recovered.png', '图 8　改为整数 100 后的恢复结果（run_id 16e746cf…15e39）', 4.65)

heading('九、文件血缘与哈希核验', 1)
add_table(['清单字段', '实测值'], [
    ['source_sha256', '0d93bcd25e28047f804238c220bdb5b1d2dadf328647c07de85ba0104e8e2854'],
    ['silver_sha256', '516b3d78aab91c1487d9a24979eaef514c5acb7309d0d82a6688a762836634dc'],
    ['gold_sha256', '654e0aeb835565d09b61f965015b61c3442c78e3cc320ca6e1e25b94e02dfc6e'],
    ['rows / contract_version', '20,000 / 2'],
    ['时间政策', '解析 UTC ISO 时间并规范化为 datetime.isoformat()'],
], widths=[1.5, 4.8], font_size=8.0)
image_with_caption('04-manifest-hash.png', '图 9　原始、Silver、Gold 三个摘要与清单一致（run_id ac8b59db…5fb408）', 5.4)
para('后续消费者先读取 silver.json，逐个重算原始文件、Silver 和 Gold 的 SHA-256，再决定是否继续计算。若原始文件变化，必须重新完成实验一往返并重跑实验二；若仅 Silver 或 Gold 损坏，则重跑实验二。不能直接改清单摘要让检查变绿。')
para('文件清单提供内容核验和最小血缘，但不是 ACID 事务，也不是完整湖仓表格式：它没有事务日志、并发提交、快照隔离、时间旅行或多文件原子提交能力。')

heading('十、必答分析题', 1)
para('1. True、"100" 与整数 100 的处理不同：契约要求 input_tokens 的类型精确为 int。True 虽可被 isinstance(True, int) 识别为整数子类，但不是业务整数；"100" 是字符串，需要上游显式修正类型。只有整数 100 同时满足类型与范围。', first_line=False)
para('2. 同 ID 内容冲突不能当作普通重复，因为两条记录对同一业务对象给出了不同事实。直接保留任意一条都会掩盖更正语义；应先定义版本、更新时间或纠错事件政策，再决定发布哪条。', first_line=False)
para('3. 原始 JSONL 与 Silver Parquet 的大小比较不是纯格式性能对照：输入还包含 100 条重复和 2 条坏记录，Silver 已清洗；JSONL 是行式文本，Parquet 采用列式编码和 Zstd 压缩。大小差异同时来自数据范围、编码和压缩。', first_line=False)
para('4. 列裁剪是只读取查询需要的列；谓词下推是把过滤条件交给底层扫描；行组是 Parquet 文件内部的一批行及其统计信息；目录分区是按日期、服务等键值拆成不同路径。本章实现了列裁剪、扫描过滤和 5 个行组，但没有实现目录分区，也不能证明所有行组都被跳过。', first_line=False)
para('5. 阈值 1200 会丢掉输入量大于 1200 的合法业务请求，本次减少 6,595 条合法独特记录和约 55.0% 的输入 Token。技术校验判断字段是否合法；业务政策判断合法请求是否属于当前分析范围。前者不能放宽，后者必须由业务目标和代价共同决定。', first_line=False)
para('6. 清单只保存文件摘要、行数、schema 和契约版本，能发现内容变化，但不能提供事务提交、并发控制、快照隔离或时间旅行。它适合作为最小血缘与完整性核验，不等于 ACID 数据库或完整湖仓表格式。', first_line=False)

heading('十一、结论与局限', 1)
para('实验完成了从原始 JSONL 到 Silver Parquet 再到 Gold 日汇总的可复现链路。基线 20,102 条记录全部有去向：20,000 条合法独特、100 条重复、2 条隔离，数量和 Token 总量均守恒；Gold 与 Silver 的请求数及 Token 总量一致。个人 1200 阈值进一步说明，合法数据也可能因业务范围被排除，必须把误拒绝代价写清楚。')
para('本实验仍有边界：所有计算在单机 CPU 完成，去重使用内存字典，数据规模约 2 万行；只有一个 Parquet 文件，没有目录分区和跨机容灾；清单不是事务系统；Gold 统计包含错误请求，只代表当前资源口径。更大规模需要外部排序、分布式去重、分区治理和并发提交机制。')

heading('十二、运行证据索引', 1)
para('运行 JSON 与同名 HTML 位于 evidence/运行记录/，原始截图位于 evidence/截图/。本报告引用的运行如下。', first_line=False)
add_table(['阶段', 'run_id', '对应截图'], [
    ['基线质量', 'e6a3a1ca90d1478688541f770a936ddf', '01-baseline-quality.png'],
    ['Gold 总量', '032c60ce145e4b5bac22eec26fa5b2af', '02-gold-totals.png'],
    ['Parquet 行组', 'f645738912114475b6f351408c814f53', '03-parquet-rowgroups.png'],
    ['清单哈希', 'ac8b59db03da425c87bfc8ddb65fb408', '04-manifest-hash.png'],
    ['政策 2000', '1e8f23cce022496591ff4d2b2dff0122', '05-policy-2000.png'],
    ['政策 1200', '0a2d4a98c7ff40e8b077a5fc97d6f2cc', '06-policy-1200.png'],
    ['契约探针', '5be251e845014e62bee30e5dbbb99ea7', '07-contract-probes.png'],
    ['预期失败', 'b54b88afad3b495f8310601935a4151b', '08-expected-failure.png'],
    ['恢复运行', '16e746cfe24e4a5f95a6eaf282c15e39', '09-recovered.png'],
], widths=[1.1, 3.0, 2.2], font_size=7.8)

heading('十三、代码仓库', 1)
para('实验一与实验二的报告、代码、运行记录和截图已汇总到以下仓库：', first_line=False)
para('https://github.com/wxzn222/cloud_computing_and_bigdata_processing', bold=True, first_line=False)

para('报告正文学号：2410311411　　实验/项目号：02', bold=True, first_line=False)

section = doc.sections[0]
footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer_p.add_run('学号 2410311411　|　实验二　数据湖与质量契约')
set_run_font(run, '宋体', 8.5)


def force_black(paragraphs):
    for paragraph in paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)


force_black(doc.paragraphs)
for table_obj in doc.tables:
    for row_obj in table_obj.rows:
        for cell_obj in row_obj.cells:
            force_black(cell_obj.paragraphs)
for section_obj in doc.sections:
    for story in (section_obj.header, section_obj.footer, section_obj.first_page_header, section_obj.first_page_footer,
                  section_obj.even_page_header, section_obj.even_page_footer):
        force_black(story.paragraphs)
        for table_obj in story.tables:
            for row_obj in table_obj.rows:
                for cell_obj in row_obj.cells:
                    force_black(cell_obj.paragraphs)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
DOCX_OUTPUTS.mkdir(parents=True, exist_ok=True)
doc.save(str(OUTPUT))
print(OUTPUT.resolve())
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables))
